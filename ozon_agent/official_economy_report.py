"""Parse official Ozon Seller 'Юнит-экономика' XLSX and apply private COGS.

This is the authoritative month-close model. Ozon supplies sales, discount
points, partner programs and all marketplace expense columns. We replace only
the Seller cost with purchase_cost * 1.35 and recalculate profit.
"""
from __future__ import annotations

import argparse
import csv
import os
from pathlib import Path

from openpyxl import load_workbook

from .business_rules import is_clearance_sku
from .unit_economics_live import load_cost_map, resolve_purchase_cost

LANDED_COST_FACTOR = 1.35

EXPENSE_COLUMNS = [
    "Вознаграждение Ozon",
    "Эквайринг",
    "Обработка отправления",
    "Логистика",
    "Доставка до места выдачи",
    "Стоимость размещения",
    "Обработка возврата",
    "Обратная логистика",
    "Утилизация",
    "Дополнительная обработка ОВХ",
    "Операционные ошибки",
    "Оплата за клик",
    "Оплата за заказ",
    "Звёздные товары",
    "Платный бренд",
    "Отзывы",
]

def _norm_header(value):
    return " ".join(str(value or "").replace("\n"," ").split())

def _base_offer(offer):
    raw=str(offer or "").strip(); up=raw.upper()
    for s in (".УЦ","-УЦ","_УЦ"," УЦ"):
        if up.endswith(s):
            return raw[:-len(s)].strip()
    return raw

def _resolve(offer,cost_map):
    cost,source,verified=resolve_purchase_cost(offer,cost_map)
    if cost is not None:
        return cost,source,verified
    if is_clearance_sku(offer):
        base=_base_offer(offer)
        if base and base!=offer:
            cost,source,verified=resolve_purchase_cost(base,cost_map)
            if cost is not None:
                return cost,"clearance_base:"+source,verified
    return None,"missing",False

def _num(value):
    try:
        return float(value or 0)
    except (TypeError,ValueError):
        return 0.0

def read_official_report(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    ws=wb.worksheets[0]
    header_row=None; headers=None
    for idx,row in enumerate(ws.iter_rows(values_only=True),start=1):
        vals=[_norm_header(v) for v in row]
        if "SKU" in vals and "Артикул" in vals and "Выручка" in vals and "Прибыль за период" in vals:
            header_row=idx; headers=vals; break
    if not header_row:
        raise RuntimeError("Ozon unit-economics header not found")
    pos={h:i for i,h in enumerate(headers) if h}
    required=["SKU","Артикул","Доставлено товаров, шт","Возвращено товаров, шт",
              "Выручка","Баллы за скидки","Программы партнёров"]
    missing=[x for x in required if x not in pos]
    if missing:
        raise RuntimeError("Missing Ozon columns: "+", ".join(missing))

    rows=[]
    for raw in ws.iter_rows(min_row=header_row+1,values_only=True):
        sku=str(raw[pos["SKU"]] or "").strip()
        offer=str(raw[pos["Артикул"]] or "").strip()
        if not sku or not offer:
            continue
        row={h:(raw[i] if i<len(raw) else None) for h,i in pos.items()}
        rows.append(row)
    return rows

def calculate(rows,cost_map,allow_unverified=True):
    out=[]
    for src in rows:
        offer=str(src.get("Артикул") or "").strip()
        sku=str(src.get("SKU") or "").strip()
        delivered=_num(src.get("Доставлено товаров, шт"))
        returned=_num(src.get("Возвращено товаров, шт"))
        net=delivered-returned
        revenue=_num(src.get("Выручка"))
        points=_num(src.get("Баллы за скидки"))
        partners=_num(src.get("Программы партнёров"))
        economic_base=revenue+points+partners
        expenses=sum(_num(src.get(col)) for col in EXPENSE_COLUMNS)

        purchase,source,verified=_resolve(offer,cost_map)
        base={
            "ozon_sku":sku,"offer_id":offer,
            "name":str(src.get("Название товара") or ""),
            "delivered_units":delivered,"returned_units":returned,"net_units":net,
            "buyer_revenue_rub":round(revenue,2),
            "ozon_discount_points_rub":round(points,2),
            "partner_programs_rub":round(partners,2),
            "economic_sales_base_rub":round(economic_base,2),
            "ozon_expenses_rub":round(expenses,2),
            "ozon_reported_profit_rub":round(_num(src.get("Прибыль за период")),2),
            "ozon_reported_unit_cost_rub":round(_num(src.get("Себестоимость")),2),
            "purchase_cost_source":source,"purchase_cost_verified":bool(verified),
        }
        if purchase is None:
            out.append({**base,"status":"MISSING_COST"});continue
        if not verified and not allow_unverified:
            out.append({**base,"status":"UNVERIFIED_COST"});continue

        landed=float(purchase)*LANDED_COST_FACTOR
        cogs=landed*net
        profit=economic_base+expenses-cogs
        margin=profit/economic_base*100 if economic_base>0 else 0.0
        status="RED" if profit<0 else ("YELLOW" if margin<20 else "GREEN")
        out.append({
            **base,
            "purchase_cost_rub":round(float(purchase),2),
            "landed_unit_cost_rub":round(landed,2),
            "cogs_rub":round(cogs,2),
            "contribution_profit_rub":round(profit,2),
            "contribution_profit_per_net_unit_rub":round(profit/net,2) if net else None,
            "contribution_margin_percent":round(margin,2),
            "status":status,
        })
    out.sort(key=lambda r:({"RED":0,"YELLOW":1,"GREEN":2}.get(r.get("status"),9),
                           float(r.get("contribution_profit_rub") or 0)))
    return out

def write_csv(path,rows):
    fields=[
        "ozon_sku","offer_id","name","delivered_units","returned_units","net_units",
        "buyer_revenue_rub","ozon_discount_points_rub","partner_programs_rub",
        "economic_sales_base_rub","ozon_expenses_rub",
        "ozon_reported_unit_cost_rub","ozon_reported_profit_rub",
        "purchase_cost_rub","landed_unit_cost_rub","cogs_rub",
        "contribution_profit_rub","contribution_profit_per_net_unit_rub",
        "contribution_margin_percent","purchase_cost_source","purchase_cost_verified","status",
    ]
    p=Path(path);p.parent.mkdir(parents=True,exist_ok=True)
    with p.open("w",encoding="utf-8-sig",newline="") as fh:
        w=csv.DictWriter(fh,fieldnames=fields,delimiter=";",extrasaction="ignore")
        w.writeheader();w.writerows(rows)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--report",required=True)
    ap.add_argument("--output",default="artifacts/private/official_ozon_profitability.csv")
    args=ap.parse_args()
    cost_map=load_cost_map()
    if not cost_map:
        raise RuntimeError("PURCHASE_COST_MAP_B64 is empty")
    rows=read_official_report(args.report)
    allow=os.getenv("ALLOW_LEGACY_COSTS","").upper() in {"YES","TRUE","1"}
    result=calculate(rows,cost_map,allow)
    write_csv(args.output,result)
    print("Official Ozon Economy parsed:",
          "rows=",len(result),
          "red=",sum(r.get("status")=="RED" for r in result),
          "yellow=",sum(r.get("status")=="YELLOW" for r in result),
          "green=",sum(r.get("status")=="GREEN" for r in result),
          "missing_cost=",sum(r.get("status")=="MISSING_COST" for r in result))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
