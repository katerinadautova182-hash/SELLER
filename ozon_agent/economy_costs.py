from collections import defaultdict
from datetime import datetime,timedelta
import io,time,requests
from openpyxl import load_workbook
from .finance_accruals import fetch_accruals_for_day,money
T_ACQ=1;T_PROMO=3;T_PROC=17;T_LAST=29;T_LOG=32;T_RET=59;T_HAND=98;T_ERR=94
RETURN_TYPES={59,45}

def dates(a,b):
 d=datetime.strptime(a,"%Y-%m-%d").date();e=datetime.strptime(b,"%Y-%m-%d").date()
 while d<=e:
  yield d.isoformat();d+=timedelta(days=1)

def services(acc):
 for p in ((acc.get("posting") or {}).get("products") or []):
  sku=str(p.get("sku") or "")
  for s in ((p.get("delivery") or {}).get("services") or []):
   try: tid=int(s.get("type_id"))
   except Exception: continue
   if sku: yield sku,tid,money(s.get("accrued"))

def placement_costs(client,a,b):
 created=client.post("/v1/report/placement/by-products/create",{"date_from":a,"date_to":b})
 code=created.get("code") or (created.get("result") or {}).get("code")
 if not code: return {}
 for _ in range(30):
  info=client.post("/v1/report/info",{"code":code})
  result=info.get("result") or {}
  url=result.get("file")
  if url:
   r=requests.get(url,timeout=60);r.raise_for_status()
   wb=load_workbook(io.BytesIO(r.content),read_only=True,data_only=True)
   out=defaultdict(float)
   for ws in wb.worksheets:
    rows=ws.iter_rows(values_only=True);header=next(rows,None)
    if not header: continue
    cols={str(v).strip():i for i,v in enumerate(header) if v is not None}
    sku_i=cols.get("SKU"); cost_i=cols.get("Начисленная стоимость размещения")
    if sku_i is None or cost_i is None: continue
    for row in rows:
     if sku_i>=len(row) or cost_i>=len(row): continue
     sku=str(row[sku_i] or "").strip()
     if sku: out[sku]-=float(row[cost_i] or 0)
   return {k:round(v,2) for k,v in out.items()}
  time.sleep(2)
 return {}

def economy_costs(client,a,b,realization):
 accr=[]
 for day in dates(a,b): accr+=fetch_accruals_for_day(client,day)
 eligible={sku:set(v.get("posting_numbers") or []) for sku,v in realization.items()}
 returns=defaultdict(set)
 for x in accr:
  unit=str(x.get("unit_number") or "")
  for sku,tid,amt in services(x):
   if sku in realization and tid in RETURN_TYPES and amt: returns[sku].add(unit)
 out=defaultdict(lambda:{"acquiring_rub":0.0,"shipment_processing_rub":0.0,"logistics_rub":0.0,"last_mile_rub":0.0,"placement_rub":0.0,"returns_rub":0.0,"operational_errors_rub":0.0,"promotion_rub":0.0})
 for x in accr:
  unit=str(x.get("unit_number") or "");cat=str(x.get("accrued_category") or "").upper()
  if cat=="POSTING":
   for sku,tid,amt in services(x):
    if sku not in realization or (unit not in eligible.get(sku,set()) and unit not in returns.get(sku,set())): continue
    if tid==T_PROC: out[sku]["shipment_processing_rub"]+=amt
    elif tid==T_LOG: out[sku]["logistics_rub"]+=amt
    elif tid in (T_LAST,T_HAND): out[sku]["last_mile_rub"]+=amt
    elif tid==T_RET: out[sku]["returns_rub"]+=amt
  elif cat=="ITEM":
   for g in ((x.get("item_fees") or {}).get("fees") or []):
    sku=str(g.get("sku") or "")
    if sku not in realization or unit not in eligible.get(sku,set()): continue
    for f in g.get("fees") or []:
     try: tid=int(f.get("type_id"))
     except Exception: continue
     amt=money(f.get("accrued"))
     if tid==T_ACQ: out[sku]["acquiring_rub"]+=amt
     elif tid==T_PROMO: out[sku]["promotion_rub"]+=amt
  elif cat=="NON_ITEM":
   f=x.get("non_item_fee") or {}
   try: tid=int(f.get("type_id"))
   except Exception: tid=-1
   if tid==T_ERR:
    for sku,posts in eligible.items():
     if unit in posts: out[sku]["operational_errors_rub"]+=money(f.get("accrued"))
 placement=placement_costs(client,a,b)
 for sku,amount in placement.items():
  if sku in realization:
   out[sku]["placement_rub"]=amount
 return {sku:{k:round(v,2) for k,v in row.items()} for sku,row in out.items()}
