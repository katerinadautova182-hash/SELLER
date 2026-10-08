"""Probe Ozon placement-by-products XLSX report for September 2026."""
import io,time,requests
from openpyxl import load_workbook
from ozon_export.client import OzonClient

def main():
    client=OzonClient()
    created=client.post("/v1/report/placement/by-products/create",{
        "date_from":"2026-09-01",
        "date_to":"2026-09-30"
    })
    code=created.get("code") or (created.get("result") or {}).get("code")
    print("PLACEMENT code:",code)
    if not code:
        raise SystemExit("No report code")
    for _ in range(30):
        info=client.post("/v1/report/info",{"code":code})
        result=info.get("result") or {}
        status=str(result.get("status") or "")
        if result.get("file"):
            r=requests.get(result["file"],timeout=60)
            r.raise_for_status()
            wb=load_workbook(io.BytesIO(r.content),read_only=True,data_only=True)
            print("PLACEMENT sheets:",wb.sheetnames)
            for ws in wb.worksheets:
                print("PLACEMENT sheet:",ws.title)
                rows=ws.iter_rows(values_only=True)
                header=next(rows,None)
                print("PLACEMENT header:",header)
                for row in rows:
                    textrow=" | ".join("" if v is None else str(v) for v in row)
                    if "MS9500" in textrow or "496958132" in textrow:
                        print("PLACEMENT MS9500:",row)
            return
        print("PLACEMENT status:",status)
        time.sleep(2)
    raise SystemExit("Placement report timeout")

if __name__=="__main__":
    main()
