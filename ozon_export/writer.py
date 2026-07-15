"""Сохранение данных в CSV (; , UTF-8-sig) и XLSX с русскими заголовками."""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from . import config


def write_table(name: str, headers: Sequence[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    """headers — список (ключ_в_словаре, заголовок_колонки_на_русском).

    Пишет одновременно output/<name>.csv и output/<name>.xlsx.
    Если rows пуст, всё равно создаёт файлы с одними заголовками —
    так пользователь видит, что раздел обработан, а не пропущен.
    """
    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    keys = [h[0] for h in headers]
    titles = [h[1] for h in headers]

    csv_path = config.OUTPUT_DIR / f"{name}.csv"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(titles)
        for row in rows:
            writer.writerow([_fmt(row.get(k, "")) for k in keys])

    xlsx_path = config.OUTPUT_DIR / f"{name}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = name[:31] or "Данные"
    ws.append(titles)
    for row in rows:
        ws.append([_fmt(row.get(k, "")) for k in keys])
    for i, title in enumerate(titles, start=1):
        width = max(len(title), 12)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 4, 60)
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)

    print(f"  -> {csv_path.name} / {xlsx_path.name}: {len(rows)} строк")


def _fmt(value: Any) -> Any:
    if isinstance(value, (list, tuple)):
        return "; ".join(str(v) for v in value)
    if isinstance(value, bool):
        return "да" if value else "нет"
    if value is None:
        return ""
    return value
